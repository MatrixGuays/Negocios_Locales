import tkinter as tk
from tkinter import ttk, messagebox, simpledialog, filedialog
import json
import os
from typing import Dict, List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class Producto:
    """Clase que representa un producto con sus talles"""
    
    def __init__(self, articulo: str, color: str, talles: Optional[Dict] = None):
        self.articulo = articulo
        self.color = color
        self.talles = talles or {str(t): {"stock": 0, "vendidos": 0} for t in range(35, 41)}
    
    def to_dict(self) -> Dict:
        return {
            "articulo": self.articulo,
            "color": self.color,
            "talles": self.talles
        }
    
    @classmethod
    def from_dict(cls, data: Dict) -> 'Producto':
        return cls(
            articulo=data.get("articulo", ""),
            color=data.get("color", ""),
            talles=data.get("talles")
        )


class InventarioManager:
    """Clase responsable de la gestión y persistencia del inventario"""
    
    def __init__(self, archivo: str):
        self.archivo = archivo
        self.marca: str = ""
        self.articulo_default: str = ""
        self.color_default: str = ""
        self.productos: List[Producto] = []
        self.cargar_datos()
    
    def cargar_datos(self):
        if os.path.exists(self.archivo):
            try:
                with open(self.archivo, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    self.marca = data.get("marca", "FRIEND")
                    self.articulo_default = data.get("articulo_default", "994")
                    self.color_default = data.get("color_default", "PLATA")
                    self.productos = [Producto.from_dict(p) for p in data.get("productos", [])]
            except Exception as e:
                print(f"Error al cargar datos: {e}")
                self._inicializar_datos_default()
        else:
            self._inicializar_datos_default()
    
    def _inicializar_datos_default(self):
        self.productos = [
            Producto("", "", {"35": {"stock": 0, "vendidos": 0}, "36": {"stock": 0, "vendidos": 0}, "37": {"stock": 0, "vendidos": 0}, "38": {"stock": 0, "vendidos": 0}, "39": {"stock": 0, "vendidos": 0}, "40": {"stock": 0, "vendidos": 0}}),
            Producto("", ""),
            Producto("", ""),
            Producto("", ""),
            Producto("", "")
        ]
    
    def guardar_datos(self):
        try:
            data = {
                "marca": self.marca,
                "articulo_default": self.articulo_default,
                "color_default": self.color_default,
                "productos": [p.to_dict() for p in self.productos]
            }
            with open(self.archivo, 'w', encoding='utf-8') as f:
                json.dump(data, f, indent=2, ensure_ascii=False)
            return True
        except Exception as e:
            print(f"Error al guardar: {e}")
            return False
    
    def exportar_excel(self, ruta: str) -> bool:
        """Exporta el inventario a un archivo Excel"""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Inventario"
            
            # Estilos
            font_titulo = Font(name='Arial', size=16, bold=True, color='FFFFFF')
            font_header = Font(name='Arial', size=11, bold=True, color='FFFFFF')
            font_normal = Font(name='Arial', size=10)
            font_bold = Font(name='Arial', size=10, bold=True)
            
            fill_marca = PatternFill(start_color='2196F3', end_color='2196F3', fill_type='solid')
            fill_header = PatternFill(start_color='9E9E9E', end_color='9E9E9E', fill_type='solid')
            fill_stock = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
            fill_vendidos = PatternFill(start_color='FFF9C4', end_color='FFF9C4', fill_type='solid')
            
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # Título (MARCA)
            ws.merge_cells('A1:H1')
            celda_titulo = ws['A1']
            celda_titulo.value = f"INVENTARIO - {self.marca.upper()}" if self.marca else "INVENTARIO"
            celda_titulo.font = font_titulo
            celda_titulo.fill = fill_marca
            celda_titulo.alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[1].height = 30
            
            # Fila de encabezados
            headers = ['ARTÍCULO', 'COLOR', '35', '36', '37', '38', '39', '40']
            for col, header in enumerate(headers, 1):
                celda = ws.cell(row=2, column=col, value=header)
                celda.font = font_header
                celda.fill = fill_header
                celda.alignment = Alignment(horizontal='center', vertical='center')
                celda.border = thin_border
            ws.row_dimensions[2].height = 25
            
            # Sub-encabezados (Stock/Vendidos)
            ws.merge_cells('A3:B3')
            ws.cell(row=3, column=1, value='STOCK / VENDIDOS').font = Font(name='Arial', size=9, italic=True)
            ws.cell(row=3, column=1).alignment = Alignment(horizontal='center')
            
            for col in range(3, 9):
                celda = ws.cell(row=3, column=col, value='Stock / Vendidos')
                celda.font = Font(name='Arial', size=9, italic=True)
                celda.alignment = Alignment(horizontal='center')
            
            # Datos de productos
            fila = 4
            for producto in self.productos:
                # Solo exportar productos con artículo o color
                if not producto.articulo and not producto.color:
                    continue
                
                # Artículo
                celda_art = ws.cell(row=fila, column=1, value=producto.articulo)
                celda_art.font = font_bold
                celda_art.alignment = Alignment(horizontal='center', vertical='center')
                celda_art.border = thin_border
                
                # Color
                celda_col = ws.cell(row=fila, column=2, value=producto.color)
                celda_col.font = font_normal
                celda_col.alignment = Alignment(horizontal='center', vertical='center')
                celda_col.border = thin_border
                
                # Talles
                for i, talle in enumerate(range(35, 41)):
                    col = i + 3
                    datos = producto.talles[str(talle)]
                    
                    # Stock en verde
                    celda_stock = ws.cell(row=fila, column=col, value=datos['stock'])
                    celda_stock.font = Font(name='Arial', size=10, bold=True, color='2E7D32')
                    celda_stock.fill = fill_stock
                    celda_stock.alignment = Alignment(horizontal='center')
                    celda_stock.border = thin_border
                    
                    # Vendidos en la fila siguiente
                    celda_vend = ws.cell(row=fila+1, column=col, value=datos['vendidos'])
                    celda_vend.font = Font(name='Arial', size=9, color='1565C0')
                    celda_vend.fill = fill_vendidos
                    celda_vend.alignment = Alignment(horizontal='center')
                    celda_vend.border = thin_border
                
                # Artículo y color también en la fila de vendidos
                ws.cell(row=fila+1, column=1).border = thin_border
                ws.cell(row=fila+1, column=2).border = thin_border
                
                fila += 2
            
            # Totales
            fila += 1
            ws.merge_cells(f'A{fila}:B{fila}')
            celda_total = ws.cell(row=fila, column=1, value='TOTALES')
            celda_total.font = Font(name='Arial', size=11, bold=True)
            celda_total.alignment = Alignment(horizontal='center')
            
            for i, talle in enumerate(range(35, 41)):
                col = i + 3
                total_stock = sum(p.talles[str(talle)]['stock'] for p in self.productos)
                total_vendidos = sum(p.talles[str(talle)]['vendidos'] for p in self.productos)
                
                celda = ws.cell(row=fila, column=col, value=f"S:{total_stock} V:{total_vendidos}")
                celda.font = Font(name='Arial', size=10, bold=True)
                celda.alignment = Alignment(horizontal='center')
                celda.border = thin_border
            
            # Ajustar ancho de columnas
            ws.column_dimensions['A'].width = 15
            ws.column_dimensions['B'].width = 20
            for col in range(3, 9):
                ws.column_dimensions[get_column_letter(col)].width = 12
            
            # Guardar
            wb.save(ruta)
            return True
        except Exception as e:
            print(f"Error al exportar: {e}")
            return False
    
    def actualizar_articulo_default(self, articulo_nuevo: str) -> bool:
        articulo_nuevo = articulo_nuevo.strip()
        if not articulo_nuevo:
            return False
        self.articulo_default = articulo_nuevo
        return self.guardar_datos()
    
    def actualizar_color_default(self, color_nuevo: str) -> bool:
        color_nuevo = color_nuevo.strip()
        if not color_nuevo:
            return False
        self.color_default = color_nuevo
        return self.guardar_datos()
    
    def actualizar_articulo(self, articulo_viejo: str, articulo_nuevo: str) -> bool:
        articulo_nuevo = articulo_nuevo.strip()
        if not articulo_nuevo:
            return False
        for producto in self.productos:
            if producto.articulo == articulo_nuevo and producto.articulo != articulo_viejo:
                return False
        for producto in self.productos:
            if producto.articulo == articulo_viejo:
                producto.articulo = articulo_nuevo
                break
        return self.guardar_datos()
    
    def actualizar_color(self, articulo: str, color_nuevo: str) -> bool:
        color_nuevo = color_nuevo.strip()
        if not color_nuevo:
            return False
        for producto in self.productos:
            if producto.articulo == articulo:
                producto.color = color_nuevo
                break
        return self.guardar_datos()
    
    def incrementar_stock(self, articulo: str, talle: str) -> bool:
        for producto in self.productos:
            if producto.articulo == articulo:
                producto.talles[talle]["stock"] += 1
                return self.guardar_datos()
        return False
    
    def decrementar_stock(self, articulo: str, talle: str) -> tuple:
        for producto in self.productos:
            if producto.articulo == articulo:
                if producto.talles[talle]["stock"] > 0:
                    producto.talles[talle]["stock"] -= 1
                    producto.talles[talle]["vendidos"] += 1
                    return True, self.guardar_datos()
                else:
                    return False, "No hay stock disponible"
        return False, "Producto no encontrado"
    
    def resetear_vendidos(self):
        for producto in self.productos:
            for talle in producto.talles:
                producto.talles[talle]["vendidos"] = 0
        return self.guardar_datos()


class InventarioApp:
    """Clase principal de la aplicación GUI"""
    
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("Inventario de Calzado - Gestión de Stock")
        self.root.geometry("1400x800")
        self.root.minsize(1200, 600)
        
        self.manager = InventarioManager("inventario.json")
        
        self._configurar_estilos()
        self.crear_interfaz()
        self.actualizar_tabla()
    
    def _configurar_estilos(self):
        style = ttk.Style()
        style.configure("Header.TLabel", font=("Arial", 10, "bold"))
        style.configure("Cell.TLabel", font=("Arial", 9))
    
    def crear_interfaz(self):
        # Frame superior con controles
        frame_superior = tk.Frame(self.root, bg="#f0f0f0")
        frame_superior.pack(fill=tk.X, padx=10, pady=10)
        
        # MARCA
        tk.Label(frame_superior, text="MARCA:", 
                font=("Arial", 12, "bold"), bg="#f0f0f0").pack(side=tk.LEFT, padx=5)
        
        self.btn_marca = tk.Button(frame_superior, 
                                text=self.manager.marca,
                                font=("Arial", 14, "bold"), 
                                fg="white", bg="#2196F3",
                                relief="raised", borderwidth=2,
                                cursor="hand2",
                                command=self.editar_marca)
        self.btn_marca.pack(side=tk.LEFT, padx=5)
        
        # ARTÍCULO
        tk.Label(frame_superior, text="ARTÍCULO:", 
                font=("Arial", 12, "bold"), bg="#f0f0f0").pack(side=tk.LEFT, padx=5)
        
        self.btn_articulo = tk.Button(frame_superior, 
                                    text=self.manager.articulo_default,
                                    font=("Arial", 14, "bold"), 
                                    fg="white", bg="#4CAF50",
                                    relief="raised", borderwidth=2,
                                    cursor="hand2",
                                    command=self.editar_articulo_default)
        self.btn_articulo.pack(side=tk.LEFT, padx=5)
        
        # COLOR
        tk.Label(frame_superior, text="COLOR:", 
                font=("Arial", 12, "bold"), bg="#f0f0f0").pack(side=tk.LEFT, padx=5)
        
        self.btn_color = tk.Button(frame_superior, 
                                text=self.manager.color_default,
                                font=("Arial", 14, "bold"), 
                                fg="white", bg="#FF9800",
                                relief="raised", borderwidth=2,
                                cursor="hand2",
                                command=self.editar_color_default)
        self.btn_color.pack(side=tk.LEFT, padx=5)
        
        # Separador
        tk.Frame(frame_superior, width=2, bg="gray").pack(side=tk.LEFT, fill=tk.Y, padx=10)
        
        # Botón EXPORTAR EXCEL (NUEVO)
        btn_exportar = tk.Button(frame_superior, text="📊 Exportar Excel", 
                            command=self.exportar_a_excel,
                            bg="#1976D2", fg="white", font=("Arial", 10, "bold"),
                            cursor="hand2", relief="raised", borderwidth=2)
        btn_exportar.pack(side=tk.LEFT, padx=5)
        
        # Botones de acción
        btn_reset = tk.Button(frame_superior, text="🔄 Resetear Vendidos", 
                            command=self.confirmar_reset,
                            bg="#FF5722", fg="white", font=("Arial", 10, "bold"),
                            cursor="hand2", relief="raised", borderwidth=2)
        btn_reset.pack(side=tk.LEFT, padx=5)
        
        btn_guardar = tk.Button(frame_superior, text="💾 Guardar", 
                            command=self.guardar_manual,
                            bg="#4CAF50", fg="white", font=("Arial", 10, "bold"),
                            cursor="hand2", relief="raised", borderwidth=2)
        btn_guardar.pack(side=tk.LEFT, padx=5)
        
        # Frame para la tabla
        frame_tabla = tk.Frame(self.root)
        frame_tabla.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        scrollbar_y = ttk.Scrollbar(frame_tabla, orient=tk.VERTICAL)
        scrollbar_x = ttk.Scrollbar(frame_tabla, orient=tk.HORIZONTAL)
        
        self.canvas = tk.Canvas(frame_tabla, 
                            yscrollcommand=scrollbar_y.set,
                            xscrollcommand=scrollbar_x.set,
                            bg="white")
        
        scrollbar_y.config(command=self.canvas.yview)
        scrollbar_x.config(command=self.canvas.xview)
        
        scrollbar_y.pack(side=tk.RIGHT, fill=tk.Y)
        scrollbar_x.pack(side=tk.BOTTOM, fill=tk.X)
        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        self.frame_interior = tk.Frame(self.canvas, bg="white")
        self.canvas_window = self.canvas.create_window((0, 0), 
                                                        window=self.frame_interior, 
                                                        anchor="nw")
        
        self.frame_interior.bind("<Configure>", 
                                lambda e: self.canvas.configure(
                                    scrollregion=self.canvas.bbox("all")))
        
        self.canvas.bind("<Configure>",
                        lambda e: self.canvas.itemconfig(self.canvas_window, width=e.width))
    
    def exportar_a_excel(self):
        """Abre diálogo para exportar a Excel"""
        ruta = filedialog.asksaveasfilename(
            title="Exportar Inventario a Excel",
            defaultextension=".xlsx",
            filetypes=[("Archivos Excel", "*.xlsx"), ("Todos los archivos", "*.*")],
            initialfile=f"inventario_{self.manager.marca}.xlsx"
        )
        
        if ruta:
            if self.manager.exportar_excel(ruta):
                messagebox.showinfo("Éxito", f"Inventario exportado correctamente a:\n{ruta}")
            else:
                messagebox.showerror("Error", "No se pudo exportar el inventario")
    
    def editar_marca(self):
        nueva_marca = simpledialog.askstring(
            "Editar Marca", 
            "Ingrese el nuevo nombre de la marca:",
            initialvalue=self.manager.marca
        )
        
        if nueva_marca and nueva_marca.strip():
            self.manager.marca = nueva_marca.strip()
            self.btn_marca.config(text=nueva_marca.strip())
            if self.manager.guardar_datos():
                messagebox.showinfo("Éxito", "Marca actualizada correctamente")
            else:
                messagebox.showerror("Error", "No se pudo guardar la marca")
    
    def editar_articulo_default(self):
        nuevo_articulo = simpledialog.askstring(
            "Editar Artículo por Defecto", 
            "Ingrese el nuevo código de artículo por defecto:",
            initialvalue=self.manager.articulo_default
        )
        
        if nuevo_articulo is None:
            return
        
        if not nuevo_articulo.strip():
            messagebox.showwarning("Validación", "El artículo no puede estar vacío")
            return
        
        if self.manager.actualizar_articulo_default(nuevo_articulo):
            self.btn_articulo.config(text=nuevo_articulo.strip())
            messagebox.showinfo("Éxito", f"Artículo por defecto actualizado a '{nuevo_articulo.strip()}'")
        else:
            messagebox.showerror("Error", "No se pudo actualizar el artículo")
    
    def editar_color_default(self):
        nuevo_color = simpledialog.askstring(
            "Editar Color por Defecto", 
            "Ingrese el nuevo color por defecto:",
            initialvalue=self.manager.color_default
        )
        
        if nuevo_color is None:
            return
        
        if not nuevo_color.strip():
            messagebox.showwarning("Validación", "El color no puede estar vacío")
            return
        
        if self.manager.actualizar_color_default(nuevo_color):
            self.btn_color.config(text=nuevo_color.strip())
            messagebox.showinfo("Éxito", f"Color por defecto actualizado a '{nuevo_color.strip()}'")
        else:
            messagebox.showerror("Error", "No se pudo actualizar el color")
    
    def editar_articulo(self, articulo_actual: str):
        nuevo_articulo = simpledialog.askstring(
            "Editar Artículo", 
            f"Ingrese el nuevo código para el artículo '{articulo_actual}':",
            initialvalue=articulo_actual
        )
        
        if nuevo_articulo is None:
            return
        
        if not nuevo_articulo.strip():
            messagebox.showwarning("Validación", "El artículo no puede estar vacío")
            return
        
        if self.manager.actualizar_articulo(articulo_actual, nuevo_articulo):
            messagebox.showinfo("Éxito", f"Artículo actualizado a '{nuevo_articulo.strip()}'")
            self.actualizar_tabla()
        else:
            messagebox.showerror("Error", "Ya existe un producto con ese código o el valor es inválido")
    
    def editar_color(self, articulo: str):
        producto_actual = None
        for p in self.manager.productos:
            if p.articulo == articulo:
                producto_actual = p
                break
        
        if not producto_actual:
            return
        
        nuevo_color = simpledialog.askstring(
            "Editar Color", 
            f"Ingrese el nuevo color para el artículo '{articulo}':",
            initialvalue=producto_actual.color
        )
        
        if nuevo_color is None:
            return
        
        if not nuevo_color.strip():
            messagebox.showwarning("Validación", "El color no puede estar vacío")
            return
        
        if self.manager.actualizar_color(articulo, nuevo_color):
            messagebox.showinfo("Éxito", f"Color actualizado a '{nuevo_color.strip()}'")
            self.actualizar_tabla()
        else:
            messagebox.showerror("Error", "No se pudo actualizar el color")
    
    def manejar_incremento(self, articulo: str, talle: str):
        if self.manager.incrementar_stock(articulo, talle):
            self.actualizar_tabla()
        else:
            messagebox.showerror("Error", "No se pudo incrementar el stock")
    
    def manejar_decremento(self, articulo: str, talle: str):
        exito, mensaje = self.manager.decrementar_stock(articulo, talle)
        
        if exito:
            self.actualizar_tabla()
        else:
            messagebox.showwarning("Stock Insuficiente", mensaje)
    
    def confirmar_reset(self):
        if messagebox.askyesno("Confirmar Reset", 
                              "¿Está seguro que desea resetear todos los contadores de vendidos a 0?\n\nEsta acción no se puede deshacer."):
            if self.manager.resetear_vendidos():
                messagebox.showinfo("Éxito", "Vendidos reseteados correctamente")
                self.actualizar_tabla()
            else:
                messagebox.showerror("Error", "No se pudo resetear los vendidos")
    
    def guardar_manual(self):
        if self.manager.guardar_datos():
            messagebox.showinfo("Éxito", "Datos guardados correctamente")
        else:
            messagebox.showerror("Error", "No se pudieron guardar los datos")
    
    def actualizar_tabla(self):
        for widget in self.frame_interior.winfo_children():
            widget.destroy()
        
        self.btn_marca.config(text=self.manager.marca)
        self.btn_articulo.config(text=self.manager.articulo_default)
        self.btn_color.config(text=self.manager.color_default)
        
        self.frame_interior.columnconfigure(0, weight=0, minsize=120)
        self.frame_interior.columnconfigure(1, weight=0, minsize=150)
        for i in range(2, 8):
            self.frame_interior.columnconfigure(i, weight=1, minsize=100)
        
        headers = ["ARTÍCULO\n(click para editar)", "COLOR\n(click para editar)"] + [str(t) for t in range(35, 41)]
        
        for col, header in enumerate(headers):
            bg_color = "#2196F3" if col == 0 else "#FFC107" if col == 1 else "#9E9E9E"
            lbl = tk.Label(self.frame_interior, text=header, 
                        font=("Arial", 9, "bold"), 
                        bg=bg_color, fg="white",
                        relief="raised", borderwidth=2,
                        padx=10, pady=10)
            lbl.grid(row=0, column=col, sticky="nsew")
        
        for fila, producto in enumerate(self.manager.productos, start=1):
            bg_color = "#f9f9f9" if fila % 2 == 0 else "white"
            
            btn_articulo = tk.Button(self.frame_interior, 
                                    text=producto.articulo,
                                    font=("Arial", 11, "bold"),
                                    bg="#E3F2FD", fg="#1565C0",
                                    relief="raised", borderwidth=2,
                                    cursor="hand2",
                                    command=lambda a=producto.articulo: self.editar_articulo(a))
            btn_articulo.grid(row=fila, column=0, sticky="nsew", padx=2, pady=2)
            
            btn_color = tk.Button(self.frame_interior, 
                                text=producto.color,
                                font=("Arial", 10),
                                bg="#FFF9C4", fg="#F57F17",
                                relief="raised", borderwidth=2,
                                cursor="hand2",
                                command=lambda a=producto.articulo: self.editar_color(a))
            btn_color.grid(row=fila, column=1, sticky="nsew", padx=2, pady=2)
            
            for i, talle in enumerate(range(35, 41)):
                col = i + 2
                talle_str = str(talle)
                datos = producto.talles[talle_str]
                
                frame_celda = tk.Frame(self.frame_interior, 
                                    bg=bg_color,
                                    relief="sunken", borderwidth=1)
                frame_celda.grid(row=fila, column=col, sticky="nsew", padx=2, pady=2)
                
                lbl_stock = tk.Label(frame_celda, 
                                    text=f"Stock: {datos['stock']}",
                                    font=("Arial", 10, "bold"),
                                    fg="#2E7D32" if datos['stock'] > 0 else "#C62828",
                                    bg=bg_color)
                lbl_stock.pack(pady=2)
                
                lbl_vendidos = tk.Label(frame_celda, 
                                    text=f"Vendidos: {datos['vendidos']}",
                                    font=("Arial", 9),
                                    fg="#1565C0",
                                    bg=bg_color)
                lbl_vendidos.pack(pady=1)
                
                frame_botones = tk.Frame(frame_celda, bg=bg_color)
                frame_botones.pack(pady=3)
                
                btn_menos = tk.Button(frame_botones, text="➖", 
                                    width=3, height=1,
                                    font=("Arial", 10, "bold"),
                                    bg="#EF5350", fg="white",
                                    relief="raised", borderwidth=2,
                                    cursor="hand2",
                                    command=lambda a=producto.articulo, t=talle_str: self.manejar_decremento(a, t))
                btn_menos.pack(side=tk.LEFT, padx=2)
                
                btn_mas = tk.Button(frame_botones, text="➕", 
                                width=3, height=1,
                                font=("Arial", 10, "bold"),
                                bg="#66BB6A", fg="white",
                                relief="raised", borderwidth=2,
                                cursor="hand2",
                                command=lambda a=producto.articulo, t=talle_str: self.manejar_incremento(a, t))
                btn_mas.pack(side=tk.LEFT, padx=2)


def main():
    root = tk.Tk()
    
    try:
        root.iconbitmap(default='')
    except:
        pass
    
    app = InventarioApp(root)
    
    def on_closing():
        if messagebox.askokcancel("Salir", "¿Desea guardar los cambios antes de salir?"):
            app.manager.guardar_datos()
        root.destroy()
    
    root.protocol("WM_DELETE_WINDOW", on_closing)
    
    root.mainloop()


if __name__ == "__main__":
    main()